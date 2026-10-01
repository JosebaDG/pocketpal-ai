import tempfile
import unittest
from pathlib import Path
import sys
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx_to_coach_roster import convert

class ConverterTests(unittest.TestCase):
    def test_underscored_heading_and_waiting_list(self):
        doc = Document()
        doc.add_paragraph('DEMO_Lunes_17:00')
        first = doc.add_table(rows=1, cols=3)
        for row in [('Código', 'Alumno', 'Email'), ('demo1', 'Lía Ejemplo', 'private@example.test')]:
            cells = first.rows[0].cells if row[0] == 'Código' else first.add_row().cells
            for cell, text in zip(cells, row): cell.text = text
        doc.add_paragraph('Lunes 17:00_Lista_De_Espera')
        waiting = doc.add_table(rows=1, cols=2)
        for cell, text in zip(waiting.rows[0].cells, ('Código', 'Alumno')): cell.text = text
        for cell, text in zip(waiting.add_row().cells, ('demo2', 'Teo Ejemplo')): cell.text = text
        doc.add_paragraph('OTRO_Jueves_18:35')
        second = doc.add_table(rows=1, cols=2)
        for cell, text in zip(second.rows[0].cells, ('Código', 'Alumno')): cell.text = text
        for cell, text in zip(second.add_row().cells, ('demo3', 'Noa Ejemplo')): cell.text = text
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'synthetic.docx'
            doc.save(path)
            result = convert(path)
        self.assertEqual(len(result['groups']), 2)
        self.assertEqual([p['status'] for p in result['participants']], ['enrolled', 'waiting', 'enrolled'])
        self.assertEqual(result['participants'][2]['groupId'], result['groups'][1]['id'])
        self.assertNotIn('private@example.test', str(result))

if __name__ == '__main__':
    unittest.main()
